            c.execute("update redemptions set status='used' where id=?", (r['id'],))
            c.commit()
    return RedirectResponse(f'/redeem-scan?code={code}', status_code=303)

@app.get('/kds', response_class=HTMLResponse)
def kds(request:Request):
    items = rows("""select oi.*, o.table_no, o.created_at from order_items oi join orders o on o.id=oi.order_id
                    where o.payment_status='unpaid' order by o.created_at asc, oi.id asc""")
    groups = {'new':[], 'accepted':[], 'done':[]}
    for it in items:
        groups.setdefault(it['kitchen_status'], []).append(it)
    return templates.TemplateResponse(request, 'kds.html', {'groups':groups})

@app.get('/prep', response_class=HTMLResponse)
def prep_page(request:Request, date: Optional[str]=None):
    target_date = date or datetime.now().strftime('%Y-%m-%d')
    menu_stats = rows("""select mi.id, oi.name, sum(oi.qty) as qty
        from order_items oi join orders o on o.id=oi.order_id join menu_items mi on mi.id=oi.menu_item_id
        where o.payment_status='paid' and substr(o.paid_at,1,10)=?
        group by mi.id, oi.name order by qty desc""", (target_date,))
    ingredient_stats = rows("""select ing.ingredient_name, ing.unit, sum(oi.qty * ing.qty_per_item) as qty
        from order_items oi join orders o on o.id=oi.order_id join ingredients ing on ing.menu_item_id=oi.menu_item_id
        where o.payment_status='paid' and substr(o.paid_at,1,10)=?
        group by ing.ingredient_name, ing.unit order by qty desc""", (target_date,))
    return templates.TemplateResponse(request, 'prep.html', {'target_date':target_date, 'menu_stats':menu_stats, 'ingredient_stats':ingredient_stats})

@app.get('/admin/ingredients', response_class=HTMLResponse)
def ingredients_page(request:Request):
    items = rows('select * from menu_items order by category,id')
    ingredients = rows("""select ing.*, mi.name as menu_name from ingredients ing join menu_items mi on mi.id=ing.menu_item_id order by mi.category, mi.id, ing.id""")
    return templates.TemplateResponse(request, 'ingredients.html', {'items':items,'ingredients':ingredients})

@app.post('/admin/ingredients/add')
def ingredient_add(menu_item_id:int=Form(...), ingredient_name:str=Form(...), qty_per_item:float=Form(1), unit:str=Form('份')):
    with conn() as c:
        c.execute('insert into ingredients(menu_item_id,ingredient_name,qty_per_item,unit) values(?,?,?,?)', (menu_item_id, ingredient_name, qty_per_item, unit)); c.commit()
    return RedirectResponse('/admin/ingredients', status_code=303)

@app.post('/admin/ingredients/{ing_id}/delete')
def ingredient_delete(ing_id:int):
    with conn() as c:
        c.execute('delete from ingredients where id=?', (ing_id,)); c.commit()
    return RedirectResponse('/admin/ingredients', status_code=303)

@app.get('/settings', response_class=HTMLResponse)
def settings_page(request:Request):
    return templates.TemplateResponse(request, 'settings.html', {'config':CONFIG})

@app.get('/sales', response_class=HTMLResponse)
def sales(request:Request, date: Optional[str] = None, message: str = '', error: str = ''):
    target_date = date or datetime.now().strftime('%Y-%m-%d')
    paid = rows("select * from orders where payment_status='paid' and substr(paid_at,1,10)=? order by paid_at desc", (target_date,))
    for o in paid:
        o['total'] = order_total(o['id'])
    stats = rows('''
        select oi.name as name, sum(oi.qty) as qty, sum(oi.qty * oi.price) as amount
        from order_items oi
        join orders o on o.id = oi.order_id
        where o.payment_status='paid' and substr(o.paid_at,1,10)=?
        group by oi.name
        order by qty desc, amount desc, name asc
    ''', (target_date,))
    total_qty = sum(int(s['qty'] or 0) for s in stats)
    total_amount = sum(int(s['amount'] or 0) for s in stats)
    return templates.TemplateResponse(request, 'sales.html', {'orders':paid,'stats':stats,'target_date':target_date,'total_qty':total_qty,'total_amount':total_amount,'message':message,'error':error})



def build_sales_excel(target_date: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    paid=rows("select * from orders where payment_status='paid' and substr(paid_at,1,10)=? order by paid_at asc",(target_date,))
    for o in paid: o['total']=order_total(o['id'])
    stats=rows("""select oi.name name,sum(oi.qty) qty,sum(oi.qty*oi.price) amount
      from order_items oi join orders o on o.id=oi.order_id
      where o.payment_status='paid' and substr(o.paid_at,1,10)=?
      group by oi.name order by qty desc,amount desc,name asc""",(target_date,))
    wb=Workbook(); ws=wb.active; ws.title="菜品銷售統計"
    ws.append([f"{target_date} 菜品銷售統計"]); ws.append(["排名","菜品","銷售數量","銷售金額"])
    for n,x in enumerate(stats,1): ws.append([n,x["name"],int(x["qty"] or 0),int(x["amount"] or 0)])
    ws.append([]); ws.append(["","合計",sum(int(x["qty"] or 0) for x in stats),sum(int(x["amount"] or 0) for x in stats)])
    ws2=wb.create_sheet("結帳紀錄"); ws2.append([f"{target_date} 結帳紀錄"])
    ws2.append(["訂單","桌號/外帶號碼","付款方式","載具","統編","金額","結帳時間"])
    for o in paid: ws2.append([f"#{o['id']}",o.get("table_no",""),o.get("payment_method",""),o.get("carrier",""),o.get("tax_id",""),int(o.get("total") or 0),o.get("paid_at","")])
    for sh in (ws,ws2):
        sh.freeze_panes="A3"
        for c in sh[1]: c.font=Font(bold=True,size=14)
        for c in sh[2]: c.font=Font(bold=True); c.alignment=Alignment(horizontal="center")
        for col in range(1,sh.max_column+1):
            letter=get_column_letter(col)
            sh.column_dimensions[letter].width=min(max(max(len(str(c.value or "")) for c in sh[letter])+3,12),32)
    b=BytesIO(); wb.save(b); b.seek(0); return b

@app.get('/sales/export')
def sales_export(date: Optional[str]=None):
    target=date or datetime.now(TAIPEI_TZ).strftime('%Y-%m-%d')
    return StreamingResponse(build_sales_excel(target),
      media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
      headers={"Content-Disposition":f'attachment; filename="{target}_sales.xlsx"'})

# V2.2：後台獨立網址 aliases
@app.get('/admin/members', response_class=HTMLResponse)
def admin_members_page(request:Request, q: str = ''):
    return members_page(request, q)

@app.get('/admin/rewards', response_class=HTMLResponse)
def admin_rewards_page(request:Request):
    return rewards_page(request)

@app.get('/admin/sales', response_class=HTMLResponse)
def admin_sales_page(request:Request, date: Optional[str] = None):
    return sales(request, date)

@app.get('/admin/prep', response_class=HTMLResponse)
def admin_prep_page(request:Request, date: Optional[str]=None):
    return prep_page(request, date)

@app.get('/admin/settings', response_class=HTMLResponse)
def admin_settings_page(request:Request):
    return settings_page(request)
@app.get('/admin/settings', response_class=HTMLResponse)
def admin_settings_page(request:Request):
    return settings_page(request)
